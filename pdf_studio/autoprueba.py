"""Autoprueba de la versión empaquetada: `PDFStudio.exe --autoprueba resultado.json`.

Comprueba, sin mostrar ventanas, que el .exe tenga todo lo que necesita: los recursos
(íconos, fuentes, diccionarios), abrir y guardar con la verificación de fidelidad, las
tareas en otro proceso (--tarea), el corrector, Word (.docx), la firma digital, OpenCV y
ocrmypdf. Escribe {prueba: "ok" | "omitida: …" | "error: …"} en el JSON y termina con 0 si todo anduvo.
No usa ningún archivo del usuario: todo se arma en una carpeta temporal.
"""
from __future__ import annotations

import json
import os
import subprocess
import sys
import tempfile
import traceback
from pathlib import Path


def _pdf(ruta: Path) -> Path:
    import pymupdf
    with pymupdf.open() as doc:
        for i in range(2):
            pag = doc.new_page()
            pag.insert_text((72, 100), f"Autoprueba de PDF Studio, página {i + 1}", fontsize=14)
            pag.draw_rect(pymupdf.Rect(72, 130, 300, 220), color=(0.2, 0.3, 0.7))
        doc.save(ruta)
    return ruta


def _recursos(tmp: Path) -> str:
    from PySide6.QtWidgets import QApplication

    from pdf_studio.config import DIR_FUENTES_INCLUIDAS, DIR_RECURSOS
    from pdf_studio.ui import iconos
    QApplication.instance() or QApplication([])
    assert (DIR_RECURSOS / "logo.ico").is_file(), "falta el logo"
    assert any(DIR_FUENTES_INCLUIDAS.glob("*.ttf")), "faltan las fuentes incluidas"
    assert not iconos.crear_icono("file-key").isNull(), "faltan los íconos"
    return "ok"


def _guardar(tmp: Path) -> str:
    from pdf_studio.core.documento import Documento
    from pdf_studio.core.paginas import RotarPagina
    documento = Documento(_pdf(tmp / "a.pdf"))
    try:
        documento.ejecutar(RotarPagina([0], 90))
        guardado = documento.guardar(tmp / "a_editado.pdf")
        assert guardado.es_fiel, "falló la verificación de fidelidad"
    finally:
        documento.cerrar()
    return "ok"


def _tarea(tmp: Path) -> str:
    from pdf_studio import procesos
    pedido = {"entrada": str(_pdf(tmp / "b.pdf")), "salida": str(tmp / "b_chico.pdf"), "nivel": "bajo"}
    (tmp / "pedido.json").write_text(json.dumps(pedido), encoding="utf-8")
    programa, argumentos = procesos.comando("comprimir", tmp / "pedido.json", tmp / "resultado.json")
    subprocess.run([programa, *argumentos], env=procesos.entorno(), timeout=180, check=True,
                   creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    salida = json.loads((tmp / "resultado.json").read_text(encoding="utf-8"))
    assert salida.get("ok"), salida.get("error", "sin resultado")
    return "ok"


def _ortografia(tmp: Path) -> str:
    from pdf_studio.core import ortografia
    assert ortografia.sospechosa("Fima") and not ortografia.sospechosa("Firma"), "el diccionario no responde"
    return "ok"


def _word(tmp: Path) -> str:
    import docx
    d = docx.Document()
    d.add_paragraph("Autoprueba")
    d.save(tmp / "c.docx")
    assert "Autoprueba" in docx.Document(tmp / "c.docx").paragraphs[0].text
    return "ok"


def _firma(tmp: Path) -> str:
    from pdf_studio.core import firma_digital as fd
    fd.crear_certificado_prueba(tmp / "c.pfx", "Autoprueba", "clave")
    fd.firmar(_pdf(tmp / "d.pdf"), tmp / "d_firmado.pdf", tmp / "c.pfx", "clave",
              fd.OpcionesFirma(visible=True, pagina=0, razon="Autoprueba"))
    [estado] = fd.validar(tmp / "d_firmado.pdf")
    assert estado.valida, estado.resumen()
    return "ok"


def _opencv(tmp: Path) -> str:
    from pdf_studio.core import escaneo
    import numpy as np
    c = escaneo.cv2()
    assert c.GaussianBlur(np.zeros((10, 10), np.uint8), (3, 3), 0).shape == (10, 10)
    return "ok"


def _ocrmypdf(tmp: Path) -> str:
    if getattr(sys, "frozen", False):
        r = subprocess.run([sys.executable, "--ocrmypdf", "--version"], capture_output=True, text=True, timeout=120,
                           creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        assert r.returncode == 0, f"código {r.returncode}"
    else:
        import ocrmypdf  # noqa: F401
    return "ok"


def _ocr_real(tmp: Path) -> str:
    """Un OCR de verdad (ocrmypdf con sus complementos y Tesseract), si Tesseract está instalado."""
    import pymupdf

    from pdf_studio.core import conversion, ocr
    try:
        conversion.herramientas_ocr()
    except Exception:
        return "omitida: Tesseract no está instalado"
    with pymupdf.open() as texto:
        texto.new_page().insert_text((72, 120), "PRUEBA DE RECONOCIMIENTO", fontsize=28)
        pix = texto[0].get_pixmap(dpi=200)
    with pymupdf.open() as escaneo:
        pag = escaneo.new_page()
        pag.insert_image(pag.rect, pixmap=pix)
        capas = ocr.reconocer(escaneo, [0], tmp, timeout=300)
        assert capas, "el OCR no devolvió ninguna capa de texto"
        ocr.AgregarCapaOCR(capas).aplicar(escaneo)
        leido = escaneo[0].get_text().upper()
    assert "PRUEBA" in leido and "RECONOCIMIENTO" in leido, "el OCR no reconoció el texto de prueba"
    return "ok"


PRUEBAS = {"recursos": _recursos, "guardar con fidelidad": _guardar, "tarea en otro proceso": _tarea,
           "corrector ortográfico": _ortografia, "Word (.docx)": _word, "firma digital": _firma,
           "OpenCV": _opencv, "ocrmypdf": _ocrmypdf, "OCR con Tesseract": _ocr_real}


def main(destino: str) -> int:
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    resultados = {}
    with tempfile.TemporaryDirectory(prefix="pdfstudio_autoprueba_") as tmp:
        for nombre, prueba in PRUEBAS.items():
            carpeta = Path(tmp) / nombre.replace(" ", "_").replace("(", "").replace(")", "").replace(".", "")
            carpeta.mkdir()
            try:
                resultados[nombre] = prueba(carpeta)
            except BaseException as e:            # noqa: BLE001 — se informa y se sigue
                detalle = traceback.format_exception_only(type(e), e)[-1].strip()
                resultados[nombre] = f"error: {detalle[:300]}"
    Path(destino).write_text(json.dumps(resultados, ensure_ascii=False, indent=2), encoding="utf-8")
    return 0 if all(v == "ok" or v.startswith("omitida") for v in resultados.values()) else 1
