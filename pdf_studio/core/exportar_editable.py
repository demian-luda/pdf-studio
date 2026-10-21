"""Exportar un documento reconstruido: .docx editable (python-docx) y PDF con texto real.

El .docx usa estilos de verdad (Normal, Título 1–3, listas con numeración de Word, tablas
con celdas combinadas, encabezado y pie con los campos de número de página), así queda
cómodo para seguir editándolo. El PDF ubica cada palabra donde estaba en el escaneo, con
la fuente elegida incluida en el archivo.

Después de escribir, se vuelve a abrir el archivo y se verifica que esté todo el texto: si
falta algo, se avisa (nunca se falla en silencio).
"""
from __future__ import annotations

import io
import re
from collections import Counter
from pathlib import Path

from pdf_studio.core.documento import ErrorDocumento, Limitacion
from pdf_studio.core.maquetado import PAGINA_X_DE_Y, texto_de, tramos
from pdf_studio.core.reconstruccion import Alineacion, Bloque, Fuente, Reconstruccion, Tipo

FUENTE_PREDETERMINADA = "Arial"
GROSOR_BORDE = 0.75


def _escribir_atomico(destino: Path, escribir) -> None:
    destino = Path(destino)
    temporal = destino.with_name(destino.name + ".pdfstudio-tmp")
    try:
        escribir(temporal)
        temporal.replace(destino)
    except PermissionError as e:
        temporal.unlink(missing_ok=True)
        raise ErrorDocumento(f"No se pudo escribir {destino}. ¿Está abierto en otro programa?") from e
    except OSError as e:
        temporal.unlink(missing_ok=True)
        raise ErrorDocumento(f"No se pudo escribir {destino}: {e}") from e


def palabras_esperadas(rec: Reconstruccion) -> Counter:
    """Las palabras que tiene que tener el documento exportado: el cuerpo de todas las páginas
    y el encabezado y el pie de la primera (en el .docx se repiten solos). Los números de
    página se reemplazan por campos, así que no cuentan."""
    esperadas: Counter = Counter()
    for n, pag in enumerate(rec.paginas):
        marco = pag.encabezado + pag.pie if n == 0 else []
        for b in pag.bloques + marco:
            en_marco = any(b is x for x in marco)
            for palabra in _palabras_de(b):
                if en_marco and palabra.isdigit():
                    continue
                esperadas[palabra] += 1
    return esperadas


def _palabras_de(b: Bloque) -> list[str]:
    if b.tipo == Tipo.IMAGEN:
        return []
    if b.tipo == Tipo.TABLA:
        return [p for c in b.celdas for x in c.bloques for p in _palabras_de(x)]
    return texto_de(b).split()


def _faltantes(esperadas: Counter, encontradas: Counter) -> list[str]:
    return sorted((esperadas - encontradas).elements())


# ---------------------------------------------------------------------- .docx
def _twips(pt: float) -> int:
    return int(round(pt * 20))


def _fuente_en(elemento_rpr, familia: str) -> None:
    """Fija la fuente en un w:rPr sacando las fuentes del tema (que tendrían prioridad)."""
    from docx.oxml.ns import qn
    rfonts = elemento_rpr.get_or_add_rFonts()
    for atributo in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(atributo), None)
    for atributo in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(atributo), familia)


def _altura_de_linea(fuente: Fuente | None) -> float:
    """Altura de una línea con interlineado sencillo en Word, en fracciones del tamaño de letra."""
    if fuente is None:
        return 1.15
    try:
        from fontTools.ttLib import TTFont
        f = TTFont(fuente.regular, fontNumber=fuente.indice_regular, lazy=True)
        os2, unidades = f["OS/2"], f["head"].unitsPerEm
        altura = (os2.usWinAscent + os2.usWinDescent) / unidades
        f.close()
        return altura if 0.8 <= altura <= 2.0 else 1.15
    except Exception:
        return 1.15


class _EscritorDocx:
    def __init__(self, rec: Reconstruccion, familia: str):
        from docx import Document
        self.rec = rec
        self.familia = familia
        self.doc = Document()
        self.altura = _altura_de_linea(rec.fuente if rec.fuente and rec.fuente.familia == familia else None)
        self._listas: dict[int, tuple[int, int, str, bool]] = {}   # contenedor -> (numId, último número, formato, numerada)
        self._siguiente_abstracto = 900
        self._siguiente_num = 900

    # -------------------------------------------------------------- estilos
    def preparar(self) -> None:
        from docx.shared import Pt, RGBColor
        rec, doc = self.rec, self.doc
        normal = doc.styles["Normal"]
        normal.font.size = Pt(rec.tamano_cuerpo)
        _fuente_en(normal.element.get_or_add_rPr(), self.familia)
        formato = normal.paragraph_format
        formato.space_before = Pt(0)
        formato.space_after = Pt(0)
        formato.line_spacing = self._multiplo(rec.interlineado_cuerpo, rec.tamano_cuerpo) or 1.0
        tamanos = {}
        for pag in rec.paginas:
            for b in pag.bloques:
                if b.tipo == Tipo.TITULO:
                    tamanos.setdefault(b.nivel, []).append(b.tamano)
        for nivel in (1, 2, 3):
            estilo = doc.styles[f"Heading {nivel}"]
            _fuente_en(estilo.element.get_or_add_rPr(), self.familia)
            valores = sorted(tamanos.get(nivel, []))
            estilo.font.size = Pt(valores[len(valores) // 2] if valores else rec.tamano_cuerpo * (1.5 - 0.15 * nivel))
            estilo.font.bold = True
            estilo.font.italic = False
            estilo.font.color.rgb = RGBColor(0, 0, 0)
            f = estilo.paragraph_format
            f.space_before = Pt(0)
            f.space_after = Pt(0)
            f.line_spacing = 1.0
            f.keep_with_next = True
        lista = doc.styles["List Paragraph"]
        lista.paragraph_format.left_indent = None
        # Página y márgenes
        s = doc.sections[0]
        primera = rec.paginas[0]
        s.page_width, s.page_height = Pt(primera.ancho), Pt(primera.alto)
        izq, arriba, der, abajo = rec.margenes
        s.left_margin, s.top_margin, s.right_margin, s.bottom_margin = Pt(izq), Pt(arriba), Pt(der), Pt(abajo)
        s.header_distance, s.footer_distance = Pt(rec.distancia_encabezado), Pt(rec.distancia_pie)
        titulo = next((texto_de(b) for pag in rec.paginas for b in pag.bloques if b.tipo == Tipo.TITULO), "")
        doc.core_properties.title = titulo[:200]
        doc.core_properties.author = ""
        doc.core_properties.comments = "Convertido desde un escaneo con PDF Studio."

    def _multiplo(self, interlineado: float, tamano: float) -> float | None:
        if interlineado <= 0 or tamano <= 0:
            return None
        return round(max(0.8, min(3.0, interlineado / (tamano * self.altura))), 2)

    @property
    def ancho_texto(self) -> float:
        izq, _, der, _ = self.rec.margenes
        return self.rec.paginas[0].ancho - izq - der

    # -------------------------------------------------------------- todo
    def escribir(self) -> None:
        self.preparar()
        cuerpo = self.doc
        primero = cuerpo.paragraphs[0] if cuerpo.paragraphs else None
        for pag in self.rec.paginas:
            for b in pag.bloques:
                self.bloque(cuerpo, b, self.rec.margenes[0], marco=False)
        if primero is not None and not primero.text and len(cuerpo.paragraphs) > 1:
            primero._p.getparent().remove(primero._p)
        s = self.doc.sections[0]
        pag = self.rec.paginas[0]
        for bloques, destino in ((pag.encabezado, s.header), (pag.pie, s.footer)):
            if not bloques:
                continue
            destino.is_linked_to_previous = False
            inicial = destino.paragraphs[0] if destino.paragraphs else None
            for b in bloques:
                self.bloque(destino, b, self.rec.margenes[0], marco=True)
            # El párrafo vacío inicial no hace falta, pero el encabezado tiene que terminar en un párrafo
            if inicial is not None and not inicial.text and len(destino._element) > 1:
                inicial._p.getparent().remove(inicial._p)
            if destino._element[-1].tag.endswith("}tbl"):
                p = destino.add_paragraph()
                self._diminuto(p)

    @staticmethod
    def _diminuto(p) -> None:
        from docx.shared import Pt
        p.paragraph_format.line_spacing = Pt(1)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)

    # -------------------------------------------------------------- bloques
    def bloque(self, contenedor, b: Bloque, izquierda: float, marco: bool, ancho: float | None = None) -> None:
        if b.tipo == Tipo.TABLA:
            if b.salto_antes:
                p = contenedor.add_paragraph()
                self._diminuto(p)
                p.paragraph_format.page_break_before = True
            elif b.espacio_antes >= 1:
                # Una tabla no tiene «espacio antes»: va como espacio después del párrafo anterior
                previo = _ultimo_parrafo(contenedor)
                if previo is not None:
                    from docx.shared import Pt
                    previo.paragraph_format.space_after = Pt(b.espacio_antes)
            self.tabla(contenedor, b, izquierda, marco)
            return
        p = contenedor.add_paragraph()
        self.formato(p, b, izquierda)
        if b.tipo == Tipo.IMAGEN:
            self.imagen(p, b, ancho or self.ancho_texto)
            return
        if b.tipo == Tipo.TITULO and not marco:
            p.style = self.doc.styles[f"Heading {max(1, min(3, b.nivel))}"]
        elif b.tipo == Tipo.ITEM:
            self.item(p, b, contenedor)
        self.texto(p, b, marco)

    def formato(self, p, b: Bloque, izquierda: float) -> None:
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
        f = p.paragraph_format
        f.alignment = {Alineacion.CENTRO: WD_ALIGN_PARAGRAPH.CENTER, Alineacion.DERECHA: WD_ALIGN_PARAGRAPH.RIGHT,
                       Alineacion.JUSTIFICADO: WD_ALIGN_PARAGRAPH.JUSTIFY}.get(b.alineacion, WD_ALIGN_PARAGRAPH.LEFT)
        if b.espacio_antes >= 1:
            f.space_before = Pt(b.espacio_antes)
        if b.salto_antes:
            f.page_break_before = True
        if b.alineacion in (Alineacion.IZQUIERDA, Alineacion.JUSTIFICADO) and b.sangria >= 1:
            f.left_indent = Pt(b.sangria)
        if b.sangria_primera and b.alineacion != Alineacion.CENTRO:
            f.first_line_indent = Pt(b.sangria_primera)
        if len(b.lineas) >= 2 and b.tipo != Tipo.TITULO:
            multiplo = self._multiplo(b.interlineado, b.tamano)
            normal = self._multiplo(self.rec.interlineado_cuerpo, self.rec.tamano_cuerpo)
            if multiplo and (normal is None or abs(multiplo - normal) > 0.08):
                f.line_spacing = multiplo

    def texto(self, p, b: Bloque, marco: bool) -> None:
        from docx.shared import Pt, RGBColor
        titulo = b.tipo == Tipo.TITULO and not marco
        tamano = b.tamano or self.rec.tamano_cuerpo
        estilo_tamano = None if titulo else self.rec.tamano_cuerpo
        texto_completo = "".join(t[0] for t in tramos(b))
        campos = marco and PAGINA_X_DE_Y.search(texto_completo)
        for texto, negrita, color, _ in tramos(b):
            partes = _partes_con_campos(texto) if campos else [(texto, None)]
            for contenido, campo in partes:
                if campo is not None:
                    run = self._campo(p, campo, contenido)
                else:
                    if not contenido:
                        continue
                    run = p.add_run(contenido)
                if not titulo or not negrita:
                    run.bold = negrita
                if estilo_tamano is None or abs(tamano - estilo_tamano) >= 0.5:
                    if not titulo:
                        run.font.size = Pt(tamano)
                if color and color != "000000":
                    run.font.color.rgb = RGBColor.from_string(color)

    def _campo(self, p, instruccion: str, valor: str):
        """Un campo de Word (PAGE o NUMPAGES) con su valor actual a la vista."""
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        run = p.add_run()
        inicio = OxmlElement("w:fldChar")
        inicio.set(qn("w:fldCharType"), "begin")
        run._r.append(inicio)
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = f" {instruccion} "
        run._r.append(instr)
        separa = OxmlElement("w:fldChar")
        separa.set(qn("w:fldCharType"), "separate")
        run._r.append(separa)
        t = OxmlElement("w:t")
        t.text = valor
        run._r.append(t)
        fin = OxmlElement("w:fldChar")
        fin.set(qn("w:fldCharType"), "end")
        run._r.append(fin)
        return run

    # -------------------------------------------------------------- listas
    def item(self, p, b: Bloque, contenedor) -> None:
        from docx.shared import Pt
        p.style = self.doc.styles["List Paragraph"]
        formato, numerada = _formato_lista(b)
        clave = id(contenedor)
        previa = self._listas.get(clave)
        if (previa is not None and previa[2] == formato and previa[3] == numerada
                and (not numerada or b.numero == previa[1] + 1 or b.nivel > 0)):
            num_id = previa[0]
        else:
            num_id = self._nueva_lista(formato, numerada, b.numero if numerada else 1)
        if b.nivel == 0:
            self._listas[clave] = (num_id, b.numero, formato, numerada)
        num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = max(0, min(2, b.nivel))
        num_pr.get_or_add_numId().val = num_id
        f = p.paragraph_format
        f.left_indent = Pt(max(b.sangria, 12))
        f.first_line_indent = Pt(min(-6.0, b.sangria_primera or -14.0))

    def _nueva_lista(self, formato: str, numerada: bool, inicio: int) -> int:
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls
        numeracion = self.doc.part.numbering_part.element
        abstracto, num = self._siguiente_abstracto, self._siguiente_num
        self._siguiente_abstracto += 1
        self._siguiente_num += 1
        niveles = []
        for nivel in range(3):
            if numerada:
                fmt, texto = _numfmt(formato, nivel)
                definicion = (f'<w:numFmt w:val="{fmt}"/><w:lvlText w:val="{texto}"/>')
                arranque = inicio if nivel == 0 else 1
            else:
                definicion = '<w:numFmt w:val="bullet"/><w:lvlText w:val="{}"/>'.format("•–◦"[nivel])
                arranque = 1
            niveles.append(f'<w:lvl w:ilvl="{nivel}"><w:start w:val="{arranque}"/>{definicion}'
                           f'<w:lvlJc w:val="left"/><w:pPr><w:ind w:left="{720 + 360 * nivel}" w:hanging="360"/>'
                           f'</w:pPr><w:rPr><w:rFonts w:ascii="{_xml(self.familia)}" w:hAnsi="{_xml(self.familia)}"/>'
                           '</w:rPr></w:lvl>')
        elemento = parse_xml(f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abstracto}">'
                             f'<w:multiLevelType w:val="hybridMultilevel"/>{"".join(niveles)}</w:abstractNum>')
        primero_num = numeracion.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num")
        if primero_num is not None:
            primero_num.addprevious(elemento)
        else:
            numeracion.append(elemento)
        numeracion.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="{num}"><w:abstractNumId w:val="{abstracto}"/>'
                                    '</w:num>'))
        return num

    # -------------------------------------------------------------- imágenes
    def imagen(self, p, b: Bloque, ancho_disponible: float) -> None:
        from docx.shared import Pt
        x0, y0, x1, y1 = b.caja
        ancho = min(x1 - x0, max(10.0, ancho_disponible))
        try:
            p.add_run().add_picture(b.archivo, width=Pt(ancho))
        except Exception as e:
            raise ErrorDocumento(f"No se pudo agregar una imagen recortada del escaneo ({type(e).__name__}).") from e

    # -------------------------------------------------------------- tablas
    def tabla(self, contenedor, b: Bloque, izquierda: float, marco: bool) -> None:
        from docx.enum.table import WD_ROW_HEIGHT_RULE
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Pt
        filas, columnas = len(b.filas) - 1, len(b.columnas) - 1
        anchos = [b.columnas[i + 1] - b.columnas[i] for i in range(columnas)]
        if hasattr(contenedor, "add_table") and contenedor is not self.doc:
            tabla = contenedor.add_table(filas, columnas, Pt(sum(anchos)))
        else:
            tabla = contenedor.add_table(rows=filas, cols=columnas)
        tabla.style = self.doc.styles["Table Grid"]
        tabla.autofit = False
        propiedades = tabla._tbl.tblPr
        disposicion = OxmlElement("w:tblLayout")
        disposicion.set(qn("w:type"), "fixed")
        propiedades.append(disposicion)
        corrimiento = b.columnas[0] - izquierda
        if abs(corrimiento) >= 2:
            sangria = OxmlElement("w:tblInd")
            sangria.set(qn("w:w"), str(_twips(corrimiento)))
            sangria.set(qn("w:type"), "dxa")
            propiedades.append(sangria)
        for i, col in enumerate(tabla._tbl.tblGrid.gridCol_lst):
            col.w = Pt(anchos[i])
        for i, fila in enumerate(tabla.rows):
            fila.height = Pt(max(8.0, b.filas[i + 1] - b.filas[i]))
            fila.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            for j, celda in enumerate(fila.cells):
                celda.width = Pt(anchos[j])
        for c in b.celdas:
            celda = tabla.cell(c.fila, c.columna)
            if c.filas > 1 or c.columnas > 1:
                celda = celda.merge(tabla.cell(c.fila + c.filas - 1, c.columna + c.columnas - 1))
            ocultos = [lado for lado, linea in (("left", c.columna), ("right", c.columna + c.columnas))
                       if linea in b.columnas_sin_borde]
            if ocultos:
                _ocultar_bordes(celda, ocultos)
            ancho_celda = sum(anchos[c.columna:c.columna + c.columnas]) - 11
            inicial = celda.paragraphs[0]
            for x in c.bloques:
                self.bloque(celda, x, b.columnas[c.columna] + 5.4, marco, ancho_celda)
            if c.bloques and not inicial.text and not inicial._p.xpath(".//w:drawing"):
                inicial._p.getparent().remove(inicial._p)
            # Una celda de Word no puede quedar sin párrafo
            if not celda.paragraphs:
                celda.add_paragraph()


def _ocultar_bordes(celda, lados: list[str]) -> None:
    """Sin borde de ese lado: la separación entre una etiqueta y su valor no estaba dibujada."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    propiedades = celda._tc.get_or_add_tcPr()
    bordes = propiedades.find(qn("w:tcBorders"))
    if bordes is None:
        bordes = OxmlElement("w:tcBorders")
        # tcBorders va después del ancho y de la combinación, antes del sombreado y la alineación
        posteriores = [propiedades.find(qn(n)) for n in ("w:shd", "w:noWrap", "w:tcMar", "w:textDirection",
                                                          "w:tcFitText", "w:vAlign", "w:hideMark")]
        posteriores = [e for e in posteriores if e is not None]
        if posteriores:
            posteriores[0].addprevious(bordes)
        else:
            propiedades.append(bordes)
    for lado in lados:
        borde = OxmlElement(f"w:{lado}")
        borde.set(qn("w:val"), "nil")
        bordes.append(borde)


def _ultimo_parrafo(contenedor):
    from docx.text.paragraph import Paragraph
    elemento = contenedor._element
    cuerpo = elemento.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}body")
    hijos = list(cuerpo if cuerpo is not None else elemento)
    hijos = [h for h in hijos if not h.tag.endswith("}sectPr")]
    if hijos and hijos[-1].tag.endswith("}p"):
        return Paragraph(hijos[-1], contenedor)
    return None


def _xml(texto: str) -> str:
    return texto.replace("&", "&amp;").replace('"', "&quot;").replace("<", "&lt;")


def _formato_lista(b: Bloque) -> tuple[str, bool]:
    if not b.numerado:
        return "•", False
    m = re.match(r"^(\()?(\d+|[a-z]+|[A-Z]+)([.)])$", b.marcador)
    if not m:
        return "1.", True
    abre, cuerpo, cierra = m.groups()
    if cuerpo.isdigit():
        clase = "1"
    elif cuerpo.lower() in ("i", "ii", "iii", "iv", "v", "vi", "vii", "viii", "ix", "x") and len(cuerpo) > 1:
        clase = "i" if cuerpo.islower() else "I"
    else:
        clase = "a" if cuerpo.islower() else "A"
    return f"{abre or ''}{clase}{cierra}", True


def _numfmt(formato: str, nivel: int) -> tuple[str, str]:
    clases = {"1": "decimal", "a": "lowerLetter", "A": "upperLetter", "i": "lowerRoman", "I": "upperRoman"}
    m = re.match(r"^(\()?([1aAiI])([.)])$", formato)
    abre, clase, cierra = m.groups() if m else ("", "1", ".")
    if nivel > 0:
        clase = "a" if nivel == 1 else "i"
    return clases[clase], f"{abre or ''}%{nivel + 1}{cierra}"


def _partes_con_campos(texto: str) -> list[tuple[str, str | None]]:
    """Separa «Página 2 de 5» en texto y campos de Word (PAGE y NUMPAGES)."""
    partes: list[tuple[str, str | None]] = []
    posicion = 0
    for m in PAGINA_X_DE_Y.finditer(texto):
        partes.append((texto[posicion:m.start(2)], None))
        partes.append((m.group(2), "PAGE"))
        if m.group(4):
            partes.append((m.group(3), None))
            partes.append((m.group(4), "NUMPAGES"))
        posicion = m.end()
    partes.append((texto[posicion:], None))
    return partes


def exportar_docx(rec: Reconstruccion, destino: Path, familia: str | None = None) -> list[str]:
    """Escribe el .docx y devuelve los avisos de la verificación (vacía si está todo)."""
    try:
        import docx  # noqa: F401
    except ImportError as e:
        raise Limitacion("Falta el paquete «python-docx». Instalalo desde el panel Faltantes.") from e
    if not rec.paginas:
        raise Limitacion("No hay nada para exportar.")
    familia = familia or (rec.fuente.familia if rec.fuente else FUENTE_PREDETERMINADA)
    escritor = _EscritorDocx(rec, familia)
    escritor.escribir()
    _escribir_atomico(Path(destino), lambda ruta: escritor.doc.save(str(ruta)))
    return verificar_docx(rec, Path(destino))


def texto_docx(ruta: Path) -> str:
    """Todo el texto del .docx (cuerpo, tablas, encabezados y pies)."""
    from docx import Document
    doc = Document(str(ruta))
    partes = []

    def recorrer(contenedor) -> None:
        for elemento in contenedor._element.iter():
            if elemento.tag.endswith("}p"):
                partes.append("".join(t.text or "" for t in elemento.iter() if t.tag.endswith("}t")))
    recorrer(doc)
    for s in doc.sections:
        recorrer(s.header)
        recorrer(s.footer)
    return "\n".join(partes)


def verificar_docx(rec: Reconstruccion, ruta: Path) -> list[str]:
    try:
        encontradas = Counter(texto_docx(ruta).split())
    except Exception as e:
        return [f"No se pudo volver a abrir el .docx para verificarlo ({type(e).__name__})."]
    faltan = _faltantes(palabras_esperadas(rec), encontradas)
    if faltan:
        return [f"En el .docx faltan {len(faltan)} palabra{'s' if len(faltan) != 1 else ''} de la reconstrucción."]
    return []


# ---------------------------------------------------------------------- PDF con texto real
def _datos_fuente(archivo: str, indice: int) -> bytes:
    datos = Path(archivo).read_bytes()
    if Path(archivo).suffix.lower() == ".ttc" or datos[:4] == b"ttcf":
        from fontTools.ttLib import TTCollection
        col = TTCollection(io.BytesIO(datos))
        buf = io.BytesIO()
        col.fonts[indice].save(buf)
        return buf.getvalue()
    return datos


def exportar_pdf(rec: Reconstruccion, destino: Path, fuente: Fuente | None = None) -> list[str]:
    """PDF con texto real: cada palabra en su lugar, con la fuente elegida incluida."""
    import pymupdf

    fuente = fuente or rec.fuente
    regular = negrita = None
    if fuente is not None:
        try:
            regular = _datos_fuente(fuente.regular, fuente.indice_regular)
            negrita = _datos_fuente(fuente.negrita or fuente.regular, fuente.indice_negrita)
        except Exception:
            regular = negrita = None
    if regular is not None:
        medidas = (pymupdf.Font(fontbuffer=regular), pymupdf.Font(fontbuffer=negrita))
    else:
        medidas = (pymupdf.Font("helv"), pymupdf.Font("hebo"))
    doc = pymupdf.open()
    try:
        for pag in rec.paginas:
            pagina = doc.new_page(width=pag.ancho, height=pag.alto)
            if regular is not None:
                pagina.insert_font(fontname="FR", fontbuffer=regular)
                pagina.insert_font(fontname="FB", fontbuffer=negrita)
                nombres = ("FR", "FB")
            else:
                nombres = ("helv", "hebo")
            for b in pag.todos():
                _dibujar(pagina, b, nombres, medidas)
        doc.subset_fonts()
        usados = {c for p in rec.palabras() for c in p.texto} | set("• ") | {
            c for pag in rec.paginas for b in pag.todos() for c in b.marcador}
        _corregir_tounicode(doc, medidas, usados)
        doc.set_metadata({"producer": "PDF Studio", "creator": "PDF Studio (convertir escaneado)",
                          "title": next((texto_de(b) for p in rec.paginas for b in p.bloques
                                         if b.tipo == Tipo.TITULO), "")[:200]})
        _escribir_atomico(Path(destino), lambda ruta: doc.save(ruta, garbage=3, deflate=True))
    finally:
        doc.close()
    return verificar_pdf(rec, Path(destino))


_ENTRADA_CMAP = re.compile(rb"(?m)^(<[0-9a-fA-F]{2,8}>[ \t]+)<([0-9a-fA-F]{4})>[ \t]*$")
_RANGO_CMAP = re.compile(rb"(?m)^<([0-9a-fA-F]{2,8})>[ \t]+<([0-9a-fA-F]{2,8})>[ \t]+<([0-9a-fA-F]{4})>[ \t]*$")


def _corregir_tounicode(doc, fuentes: tuple, usados: set[str]) -> None:
    """Cuando dos caracteres comparten el glifo (el guion y el guion blando, «(» y «﴾»), MuPDF
    anota en la tabla ToUnicode el que no es: al copiar el texto saldría otro carácter. Cada
    entrada que apunta a un carácter que no se usó se cambia por el usado con el mismo glifo."""
    glifos = []
    for f in fuentes:
        propios: dict[int, str] = {}
        for c in sorted(usados):
            g = f.has_glyph(ord(c))
            if g:
                propios.setdefault(g, c)
        glifos.append((f, propios))

    def correcto(u: int) -> int:
        if chr(u) in usados:
            return u
        for f, propios in glifos:
            c = propios.get(f.has_glyph(u))
            if c is not None and len(c.encode("utf-16-be")) == 2:
                return ord(c)
        return u

    def corregir(m: re.Match) -> bytes:
        u = int(m.group(2), 16)
        nuevo = correcto(u)
        return m.group(0) if nuevo == u else m.group(1) + f"<{nuevo:04x}>".encode()

    def corregir_rango(m: re.Match) -> bytes:
        # <a> <b> <u>: los códigos a..b van a u, u+1, ... Si alguno está mal, se parte el rango
        a, b, u = int(m.group(1), 16), int(m.group(2), 16), int(m.group(3), 16)
        if b < a or b - a > 255:
            return m.group(0)
        nuevos = [correcto(u + i) for i in range(b - a + 1)]
        if nuevos == [u + i for i in range(b - a + 1)]:
            return m.group(0)
        ancho = len(m.group(1))
        return b"\n".join(f"<{a + i:0{ancho}x}> <{a + i:0{ancho}x}> <{v:04x}>".encode() for i, v in enumerate(nuevos))

    for xref in range(1, doc.xref_length()):
        try:
            clave = doc.xref_get_key(xref, "ToUnicode")
        except Exception:
            continue
        if clave[0] != "xref":
            continue
        tabla = int(clave[1].split()[0])
        datos = doc.xref_stream(tabla)
        if datos:
            nuevo = _RANGO_CMAP.sub(corregir_rango, _ENTRADA_CMAP.sub(corregir, datos))
            if nuevo != datos:
                doc.update_stream(tabla, nuevo)


def _dibujar(pagina, b: Bloque, nombres: tuple[str, str], medidas: tuple) -> None:
    import pymupdf
    if b.tipo == Tipo.IMAGEN:
        pagina.insert_image(pymupdf.Rect(b.caja), filename=b.archivo, keep_proportion=False)
        return
    if b.tipo == Tipo.TABLA:
        forma = pagina.new_shape()
        for c in b.celdas:
            x0, x1 = b.columnas[c.columna], b.columnas[c.columna + c.columnas]
            y0, y1 = b.filas[c.fila], b.filas[c.fila + c.filas]
            forma.draw_line((x0, y0), (x1, y0))
            forma.draw_line((x0, y1), (x1, y1))
            if c.columna not in b.columnas_sin_borde:
                forma.draw_line((x0, y0), (x0, y1))
            if c.columna + c.columnas not in b.columnas_sin_borde:
                forma.draw_line((x1, y0), (x1, y1))
        forma.finish(color=(0, 0, 0), width=GROSOR_BORDE)
        forma.commit()
        for c in b.celdas:
            for x in c.bloques:
                _dibujar(pagina, x, nombres, medidas)
        return
    if b.marcador and b.lineas:
        primera = b.lineas[0]
        marcador = "•" if not b.numerado else b.marcador
        # Con un espacio después: si no, al copiar el texto el marcador se pega a la primera palabra
        pagina.insert_text((b.x_marcador, primera.base), marcador + " ", fontname=nombres[0],
                           fontsize=primera.tamano)
    for linea in b.lineas:
        for i, p in enumerate(linea.palabras):
            color = tuple(int(p.color[k:k + 2], 16) / 255 for k in (0, 2, 4)) if p.color else (0, 0, 0)
            negrita = 1 if p.negrita else 0
            origen = pymupdf.Point(p.caja[0], linea.base)
            # Si la palabra no entra hasta la siguiente (la fuente es un poco más ancha que la del
            # escaneo, o el texto se corrigió), se angosta: pisarla las uniría al copiar el texto
            natural = medidas[negrita].text_length(p.texto, fontsize=linea.tamano)
            siguiente = linea.palabras[i + 1].caja[0] if i + 1 < len(linea.palabras) else None
            lugar = (siguiente - p.caja[0] - 0.08 * linea.tamano if siguiente is not None
                     else (p.caja[2] - p.caja[0]) * 1.25 + 0.5 * linea.tamano)
            escala = max(0.25, min(1.0, lugar / natural)) if natural > 0 else 1.0
            # Con un espacio de verdad después, las palabras quedan separadas al copiar el texto
            # aunque el hueco sea chico (también la última: otro renglón puede seguir a la misma altura)
            texto = p.texto + " "
            pagina.insert_text(origen, texto, fontname=nombres[negrita], fontsize=linea.tamano, color=color,
                               morph=(origen, pymupdf.Matrix(escala, 1)) if escala < 0.999 else None)


def verificar_pdf(rec: Reconstruccion, ruta: Path) -> list[str]:
    import pymupdf
    try:
        with pymupdf.open(ruta) as doc:
            encontradas = Counter(w[4] for pag in doc for w in pag.get_text("words"))
    except Exception as e:
        return [f"No se pudo volver a abrir el PDF para verificarlo ({type(e).__name__})."]
    esperadas: Counter = Counter()
    for pag in rec.paginas:
        for b in pag.todos():
            for p in b.palabras:
                for parte in p.texto.split():
                    esperadas[parte] += 1
    faltan = _faltantes(esperadas, encontradas)
    if faltan:
        return [f"En el PDF faltan {len(faltan)} palabra{'s' if len(faltan) != 1 else ''} de la reconstrucción."]
    return []
