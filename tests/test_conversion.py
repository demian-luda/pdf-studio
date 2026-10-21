"""Convertir escaneado a documento editable: maquetado (renglones, párrafos, listas, títulos,
encabezado y pie), tablas y regiones de imagen con OpenCV, exportación a .docx y PDF con
texto real, la pantalla de revisión y, con Tesseract, SOPs sintéticos escaneados de punta
a punta."""
from __future__ import annotations

import difflib
import json
import re
from collections import Counter
from pathlib import Path

import numpy as np
import pymupdf
import pytest

pytest.importorskip("cv2")
pytest.importorskip("docx")

import cv2  # noqa: E402

from pdf_studio.core import conversion, maquetado  # noqa: E402
from pdf_studio.core.exportar_editable import (_partes_con_campos, exportar_docx, exportar_pdf,  # noqa: E402
                                               texto_docx)
from pdf_studio.core.maquetado import (AnalisisPagina, Grilla, PalabraOCR, RegionImagen, maquetar,  # noqa: E402
                                       texto_de)
from pdf_studio.core.reconstruccion import (Alineacion, Palabra, Reconstruccion, Tipo)  # noqa: E402
from tests import sops_sinteticos as sop  # noqa: E402
from tests.conftest import hash_archivo  # noqa: E402


def _hay_tesseract() -> bool:
    try:
        conversion.herramientas_ocr()
        return True
    except Exception:
        return False


requiere_tesseract = pytest.mark.skipif(not _hay_tesseract(), reason="Tesseract con español no está instalado")


# ---------------------------------------------------------------------- palabras armadas a mano
class Hoja:
    """Arma palabras como las daría el OCR: renglones con su altura, en puntos."""

    def __init__(self, numero: int = 0, ancho: float = 595.0, alto: float = 842.0):
        self.a = AnalisisPagina(numero, ancho, alto)
        self._n = 0
        self._linea = 0

    def renglon(self, texto: str, x: float, base: float, tamano: float = 10.0, negrita: bool = False,
                conf: float = 95.0, ancho_letra: float = 0.5) -> float:
        """Escribe el renglón y devuelve dónde termina (x)."""
        self._linea += 1
        for palabra in texto.split():
            ancho = len(palabra) * tamano * ancho_letra
            clase = maquetado._clase_palabra(palabra) or "x"
            base_clase, _, pata = clase.partition("+")
            arriba = maquetado.FACTORES[base_clase] * tamano
            abajo = maquetado.FACTORES["pata"] * tamano if pata else 0.0
            p = Palabra(f"p{self.a.numero}-{self._n}", palabra, (x, base - arriba, x + ancho, base + abajo), conf,
                        trazo=(0.15 if negrita else 0.09) * tamano)
            self._n += 1
            self.a.palabras.append(PalabraOCR(p, (self._linea, 0, 0)))
            x += ancho + 0.3 * tamano
        return x


def _parrafo(h: Hoja, texto: str, y: float, ancho: float = 400.0, x: float = 72.0, tamano: float = 10.0) -> float:
    """Reparte el texto en renglones de como mucho `ancho` (como un procesador de texto)."""
    renglon = []
    for palabra in texto.split():
        prueba = renglon + [palabra]
        if renglon and len(" ".join(prueba)) * tamano * 0.5 > ancho:
            h.renglon(" ".join(renglon), x, y, tamano)
            y += tamano * 1.3
            renglon = [palabra]
        else:
            renglon = prueba
    if renglon:
        h.renglon(" ".join(renglon), x, y, tamano)
        y += tamano * 1.3
    return y


LARGO = ("Este procedimiento se aplica a todos los equipos de uso compartido del sector incluidas las balanzas "
         "analíticas los agitadores magnéticos y las campanas de flujo laminar del laboratorio central")


def test_parrafo_de_varios_renglones_se_une():
    h = Hoja()
    _parrafo(h, LARGO, 100)
    rec = maquetar([h.a])
    parrafos = [b for b in rec.paginas[0].bloques if b.tipo == Tipo.PARRAFO]
    assert len(parrafos) == 1 and len(parrafos[0].lineas) >= 3
    assert texto_de(parrafos[0]) == LARGO


def test_renglon_corto_corta_el_parrafo():
    h = Hoja()
    y = _parrafo(h, LARGO, 100)
    # El último renglón de LARGO es corto: ahí terminó el párrafo, aunque no haya espacio extra
    h.renglon("Fin.", 72, y)
    _parrafo(h, "Otro párrafo que empieza en el renglón siguiente sin espacio extra entre ambos", y + 13)
    rec = maquetar([h.a])
    textos = [texto_de(b) for b in rec.paginas[0].bloques]
    assert textos[0] == LARGO and textos[1] == "Fin." and textos[2].startswith("Otro párrafo")


def test_guion_de_corte_se_pega():
    h = Hoja()
    h.renglon("La limpieza incluye la desin-", 72, 100, ancho_letra=1.6)
    h.renglon("fección de todas las superficies", 72, 113)
    rec = maquetar([h.a])
    assert texto_de(rec.paginas[0].bloques[0]) == "La limpieza incluye la desinfección de todas las superficies"


def test_listas_con_vinetas_y_numeradas():
    h = Hoja()
    h.renglon("Responsables del área:", 72, 100)
    for i, item in enumerate(["El analista limpia el equipo", "El supervisor firma la planilla"]):
        h.renglon(f"• {item}", 90, 115 + 15 * i)
    for i, item in enumerate(["Apagar el equipo", "Retirar los restos", "Registrar la fecha"], 1):
        h.renglon(f"{i}. {item}", 90, 160 + 15 * i)
    rec = maquetar([h.a])
    items = [b for b in rec.paginas[0].bloques if b.tipo == Tipo.ITEM]
    vinetas, numerados = [b for b in items if not b.numerado], [b for b in items if b.numerado]
    assert [texto_de(b) for b in vinetas] == ["El analista limpia el equipo", "El supervisor firma la planilla"]
    assert [b.numero for b in numerados] == [1, 2, 3] and texto_de(numerados[1]) == "Retirar los restos"
    assert all(b.marcador == f"{n}." for n, b in enumerate(numerados, 1))


def test_vineta_dibujada_se_reconoce():
    h = Hoja()
    h.renglon("Usar guantes de nitrilo", 100, 100)
    h.a.vinetas.append((86.0, 94.0, 90.0, 98.0))       # un punto a la izquierda del renglón
    rec = maquetar([h.a])
    b = rec.paginas[0].bloques[0]
    assert b.tipo == Tipo.ITEM and not b.numerado and texto_de(b) == "Usar guantes de nitrilo"


def test_titulos_por_tamano_y_negrita():
    h = Hoja()
    h.renglon("INFORME DE LIMPIEZA", 200, 80, tamano=16, negrita=True)
    h.renglon("1. OBJETIVO", 72, 120, tamano=12, negrita=True)
    _parrafo(h, LARGO, 140)
    h.renglon("2. ALCANCE", 72, 220, tamano=12, negrita=True)
    _parrafo(h, LARGO, 240)
    rec = maquetar([h.a])
    titulos = [b for b in rec.paginas[0].bloques if b.tipo == Tipo.TITULO]
    assert [(texto_de(b), b.nivel) for b in titulos] == [("INFORME DE LIMPIEZA", 1), ("1. OBJETIVO", 2),
                                                         ("2. ALCANCE", 2)]
    assert all(b.negrita for b in titulos)
    # «1.» de un título en negrita no es una lista
    assert not any(b.tipo == Tipo.ITEM for b in rec.paginas[0].bloques)


def test_centrado_y_justificado():
    h = Hoja()
    h.renglon("Título centrado", 260, 80)
    _parrafo(h, LARGO, 120, ancho=451)
    rec = maquetar([h.a])
    assert rec.paginas[0].bloques[0].alineacion == Alineacion.CENTRO


def test_encabezado_y_pie_que_se_repiten():
    paginas = []
    for n in range(3):
        h = Hoja(n)
        h.renglon("SOP-QC-014 Versión 03", 72, 40, tamano=9)
        _parrafo(h, LARGO + f" página {n}", 150)
        h.renglon(f"Página {n + 1} de 3", 260, 810, tamano=9)
        paginas.append(h.a)
    rec = maquetar(paginas)
    for pag in rec.paginas:
        assert [texto_de(b) for b in pag.encabezado] == ["SOP-QC-014 Versión 03"]
        assert texto_de(pag.pie[0]).startswith("Página")
        assert all("Página" not in texto_de(b) for b in pag.bloques)
    assert rec.distancia_encabezado < 40 and rec.margenes[1] > 45
    assert rec.paginas[1].bloques[0].salto_antes and not rec.paginas[0].bloques[0].salto_antes


def test_pie_de_una_sola_pagina():
    h = Hoja()
    _parrafo(h, LARGO, 150)
    h.renglon("Página 1 de 1", 260, 810, tamano=9)
    rec = maquetar([h.a])
    assert texto_de(rec.paginas[0].pie[0]) == "Página 1 de 1"


def test_tabla_con_celdas_combinadas_e_imagen():
    h = Hoja()
    # Más abajo de la franja del encabezado (si no, con «Código» sería el encabezado)
    grilla = Grilla([300.0, 330.0, 360.0], [72.0, 200.0, 400.0], [(0, 0, 2, 1), (0, 1, 1, 1), (1, 1, 1, 1)])
    h.a.grillas.append(grilla)
    h.renglon("Código: SOP-1", 210, 320)
    h.renglon("Versión: 2", 210, 350)
    h.a.imagenes.append(RegionImagen((90.0, 305.0, 180.0, 355.0), "logo.png", "logo"))
    _parrafo(h, LARGO, 400)
    rec = maquetar([h.a])
    tabla = next(b for b in rec.paginas[0].bloques if b.tipo == Tipo.TABLA)
    celdas = {(c.fila, c.columna): c for c in tabla.celdas}
    assert celdas[(0, 0)].filas == 2 and celdas[(0, 0)].bloques[0].tipo == Tipo.IMAGEN
    assert texto_de(celdas[(0, 1)].bloques[0]) == "Código: SOP-1"
    assert texto_de(celdas[(1, 1)].bloques[0]) == "Versión: 2"
    # Ni el texto ni la imagen de la tabla aparecen sueltos
    assert [b.tipo for b in rec.paginas[0].bloques] == [Tipo.TABLA, Tipo.PARRAFO]


def test_tamano_desde_la_altura_de_las_palabras():
    p = Palabra("x", "Hola", (0, 0, 30, 7.2), 95)          # mayúscula y letras altas: 0,72 del tamaño
    assert abs(maquetado.tamano_palabra(p) - 10) < 0.1
    p = Palabra("x", "pega", (0, 0, 30, 7.3), 95)          # «x» + patas
    assert abs(maquetado.tamano_palabra(p) - 10) < 0.2
    assert maquetado.tamano_palabra(Palabra("x", ".", (0, 0, 2, 2), 95)) is None
    con_exceso = {**maquetado.FACTORES, "exceso": 0.72}
    assert abs(maquetado.tamano_palabra(Palabra("x", "Hola", (0, 0, 30, 7.92), 95), con_exceso) - 10) < 0.1


def test_roundtrip_del_modelo():
    h = Hoja()
    h.renglon("1. OBJETIVO", 72, 80, tamano=12, negrita=True)
    _parrafo(h, LARGO, 100)
    h.renglon("• Un ítem", 90, 160)
    rec = maquetar([h.a])
    copia = Reconstruccion.desde_dict(json.loads(json.dumps(rec.a_dict())))
    assert copia.a_dict() == rec.a_dict()


def test_partes_con_campos_de_pagina():
    assert _partes_con_campos("Página 2 de 5") == [("Página ", None), ("2", "PAGE"), (" de ", None),
                                                   ("5", "NUMPAGES"), ("", None)]
    assert _partes_con_campos("Pág. 3") == [("Pág. ", None), ("3", "PAGE"), ("", None)]


def test_leer_tsv():
    tsv = ("level\tpage_num\tblock_num\tpar_num\tline_num\tword_num\tleft\ttop\twidth\theight\tconf\ttext\n"
           "4\t1\t1\t1\t1\t0\t10\t10\t200\t20\t-1\t\n"
           "5\t1\t1\t1\t1\t1\t10\t10\t50\t20\t95.5\tHola\n"
           "5\t1\t1\t1\t1\t2\t70\t10\t60\t20\t41.0\tmundo\n"
           "5\t1\t1\t1\t1\t3\t140\t10\t5\t20\t95.0\t \n")
    palabras = conversion.leer_tsv(tsv)
    assert [(p.texto, p.conf, p.caja) for p in palabras] == [("Hola", 95.5, (10, 10, 60, 30)),
                                                             ("mundo", 41.0, (70, 10, 130, 30))]


# ---------------------------------------------------------------------- imagen: tablas y regiones
def _lienzo(w: int = 1600, h: int = 1200) -> np.ndarray:
    return np.zeros((h, w), np.uint8)


def test_grilla_con_celda_combinada():
    tinta = _lienzo()
    for y in (100, 200, 300):
        cv2.line(tinta, (100, y), (1300, y), 255, 3)
    cv2.line(tinta, (100, 200), (500, 200), 0, 7)               # la primera columna no tiene esa línea
    for x in (100, 500, 900, 1300):
        cv2.line(tinta, (x, 100), (x, 300), 255, 3)
    grillas, mascara = conversion.detectar_grillas(tinta)
    assert len(grillas) == 1
    g = grillas[0]
    assert len(g.filas) == 3 and len(g.columnas) == 4
    assert (0, 0, 2, 1) in g.celdas and len(g.celdas) == 5
    assert mascara[200, 700] > 0 and mascara[50, 50] == 0


def test_trazo_de_firma_no_agrega_filas():
    tinta = _lienzo()
    for y in (100, 300):
        cv2.line(tinta, (100, y), (1300, y), 255, 3)
    for x in (100, 700, 1300):
        cv2.line(tinta, (x, 100), (x, 300), 255, 3)
    cv2.line(tinta, (200, 220), (600, 225), 255, 3)             # firma con un trazo casi recto
    grillas, _ = conversion.detectar_grillas(tinta)
    assert len(grillas) == 1 and len(grillas[0].filas) == 2 and len(grillas[0].celdas) == 2


def test_logo_relleno_no_es_una_tabla():
    tinta = _lienzo()
    cv2.circle(tinta, (400, 400), 150, 255, -1)
    cv2.rectangle(tinta, (800, 300), (1100, 500), 255, -1)
    grillas, _ = conversion.detectar_grillas(tinta)
    assert grillas == []


def test_regiones_firma_sello_y_vinetas():
    h, w = 1400, 1200
    gris = np.full((h, w), 255, np.uint8)
    color = np.full((h, w, 3), 255, np.uint8)
    palabras = []
    for i in range(6):                                           # renglones de texto «reconocido»
        y = 100 + i * 60
        cv2.rectangle(gris, (200, y), (900, y + 30), 0, 2)
        palabras.append(conversion.PalabraTSV("texto", (198, y - 2, 902, y + 32), 95, (1, 1, i)))
        cv2.circle(gris, (160, y + 15), 7, 0, -1)                # viñeta
    xs = np.linspace(200, 700, 80)
    puntos = np.stack([xs, 700 + 40 * np.sin(xs / 25)], axis=1).astype(np.int32)
    cv2.polylines(gris, [puntos], False, 0, 3)                    # firma
    cv2.circle(color, (900, 1100), 120, (200, 60, 30), 8)        # sello azul
    color[gris < 60] = 30                                       # el texto y la firma, negros
    _, tinta = cv2.threshold(gris, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)
    regiones, vinetas = conversion.regiones_no_texto(tinta, color, palabras, 32.0, [])
    clases = sorted(r.clase for r in regiones)
    assert len(regiones) == 2 and "firma" in clases and "sello" in clases
    assert len(vinetas) == 6


def test_grosor_de_trazo():
    fino, grueso = _lienzo(400, 200), _lienzo(400, 200)
    cv2.line(fino, (20, 100), (380, 100), 255, 3)
    cv2.line(grueso, (20, 100), (380, 100), 255, 7)
    a = conversion.grosor_trazo(fino, (0, 0, 400, 200))
    b = conversion.grosor_trazo(grueso, (0, 0, 400, 200))
    assert 2 < a < 5 and 6 < b < 9


# ---------------------------------------------------------------------- exportación
@pytest.fixture
def rec_armada(tmp_path):
    """Una reconstrucción armada a mano: título, párrafo, listas, tabla con imagen, encabezado y pie."""
    from PIL import Image
    logo = tmp_path / "logo.png"
    Image.new("RGB", (60, 40), (30, 120, 60)).save(logo)
    paginas = []
    for n in range(2):
        h = Hoja(n)
        h.renglon("SOP-QC-014 Versión 03", 72, 40, tamano=9)
        if n == 0:
            h.renglon("PROCEDIMIENTO DE LIMPIEZA", 200, 90, tamano=16, negrita=True)
            h.renglon("1. OBJETIVO", 72, 125, tamano=12, negrita=True)
            y = _parrafo(h, LARGO, 145)
            for i, item in enumerate(["El analista limpia", "El supervisor firma"]):
                h.renglon(f"• {item}", 90, y + 10 + 14 * i)
            for i, item in enumerate(["Apagar el equipo", "Registrar la fecha"], 1):
                h.renglon(f"{i}) {item}", 90, y + 50 + 14 * i)
        else:
            h.a.grillas.append(Grilla([100.0, 125.0, 175.0], [72.0, 250.0, 520.0],
                                      [(0, 0, 1, 1), (0, 1, 1, 1), (1, 0, 1, 2)]))
            h.renglon("Material", 80, 118, negrita=True)
            h.renglon("Cantidad", 260, 118, negrita=True)
            h.a.imagenes.append(RegionImagen((120.0, 130.0, 180.0, 170.0), str(logo), "firma"))
            _parrafo(h, "Copia controlada del documento", 220)
        h.renglon(f"Página {n + 1} de 2", 260, 810, tamano=9)
        paginas.append(h.a)
    return maquetar(paginas)


def test_exportar_docx_estructura(rec_armada, tmp_path):
    from docx import Document
    destino = tmp_path / "salida.docx"
    assert exportar_docx(rec_armada, destino, "Arial") == []          # la verificación no encontró faltantes
    doc = Document(str(destino))
    estilos = [(p.style.name, p.text) for p in doc.paragraphs if p.text]
    assert ("Heading 1", "PROCEDIMIENTO DE LIMPIEZA") in estilos and ("Heading 2", "1. OBJETIVO") in estilos
    items = [p for p in doc.paragraphs if p._p.pPr is not None and p._p.pPr.numPr is not None]
    assert [p.text for p in items] == ["El analista limpia", "El supervisor firma", "Apagar el equipo",
                                       "Registrar la fecha"]
    # Las viñetas y la lista numerada son listas distintas de Word
    assert len({p._p.pPr.numPr.numId.val for p in items}) == 2
    numeracion = doc.part.numbering_part.element.xml
    assert 'w:val="bullet"' in numeracion and 'w:lvlText w:val="%1)"' in numeracion
    tabla = doc.tables[0]
    assert len(tabla.rows) == 2 and tabla.cell(0, 0).text == "Material"
    assert tabla.cell(1, 0)._tc is tabla.cell(1, 1)._tc                 # celda combinada
    assert tabla.cell(1, 0)._tc.xpath(".//w:drawing")                    # con la imagen adentro
    s = doc.sections[0]
    assert abs(s.page_width.pt - 595) < 1 and abs(s.page_height.pt - 842) < 1
    assert "SOP-QC-014" in s.header.paragraphs[0].text
    pie = s.footer._element.xml
    assert "PAGE" in pie and "NUMPAGES" in pie
    # Lo de la segunda página empieza en una página nueva
    assert any(p.paragraph_format.page_break_before for p in doc.paragraphs)
    assert doc.styles["Normal"].font.name == "Arial"


def test_exportar_docx_avisa_si_falta_texto(rec_armada, tmp_path, monkeypatch):
    from pdf_studio.core import exportar_editable
    destino = tmp_path / "salida.docx"
    original = exportar_editable.texto_docx
    monkeypatch.setattr(exportar_editable, "texto_docx", lambda ruta: original(ruta).replace("analista", ""))
    avisos = exportar_docx(rec_armada, destino, "Arial")
    assert avisos and "faltan 1 palabra" in avisos[0]


def test_exportar_docx_usa_las_correcciones(rec_armada, tmp_path):
    palabra = next(p for p in rec_armada.palabras() if p.texto == "analista")
    palabra.texto, palabra.revisada = "técnico", True
    destino = tmp_path / "salida.docx"
    assert exportar_docx(rec_armada, destino, "Arial") == []
    texto = texto_docx(destino)
    assert "El técnico limpia" in texto and "analista" not in texto


def test_exportar_pdf_con_texto_real(rec_armada, tmp_path):
    rec_armada.fuente = next((f for f in conversion.fuentes_disponibles() if f.clase == "sans"), None)
    destino = tmp_path / "salida.pdf"
    assert exportar_pdf(rec_armada, destino) == []
    with pymupdf.open(destino) as doc:
        assert doc.page_count == 2
        texto = doc[0].get_text()
        assert "SOP-QC-014" in texto and "\xad" not in texto          # guiones comunes, no blandos
        assert "PROCEDIMIENTO DE LIMPIEZA" in texto
        if rec_armada.fuente is not None:
            assert any(rec_armada.fuente.familia.replace(" ", "") in f[3].replace(" ", "")
                       for f in doc.get_page_fonts(0))
        assert doc[1].get_images()                                     # la imagen recortada


def test_pdf_copia_el_texto_tal_cual_con_todas_las_fuentes(tmp_path):
    """Caracteres que comparten glifo con otros (guion, «;», paréntesis), palabras pegadas,
    renglones a la misma altura y una palabra dudosa más ancha que su caja: el texto que se
    copia del PDF tiene que ser exactamente el reconstruido."""
    h = Hoja()
    h.renglon("SOP-QC-014 (12) (a) x; pre-aprobado", 72, 100)
    fin = h.renglon("y a", 72, 130)
    h.renglon("— sigue", fin + 0.5, 130)                         # otro renglón del OCR, pegado
    h.renglon("palabra", 72, 160, ancho_letra=0.2)                # caja muy angosta
    h.renglon("siguiente", 72 + 7 * 10 * 0.2 + 2, 160)
    rec = maquetar([h.a])
    esperadas = Counter(w for p in rec.palabras() for w in p.texto.split())
    for fuente in conversion.fuentes_disponibles():
        rec.fuente = fuente
        destino = tmp_path / f"{fuente.familia}.pdf"
        assert exportar_pdf(rec, destino) == [], fuente.familia
        with pymupdf.open(destino) as doc:
            encontradas = Counter(w[4] for w in doc[0].get_text("words"))
            texto = doc[0].get_text()
        assert encontradas == esperadas, fuente.familia
        assert " " not in texto and "­" not in texto, fuente.familia


def test_exportar_no_deja_temporales_si_falla(rec_armada, tmp_path):
    from pdf_studio.core.documento import ErrorDocumento
    carpeta = tmp_path / "no_existe"
    with pytest.raises(ErrorDocumento):
        exportar_docx(rec_armada, carpeta / "x.docx", "Arial")
    assert not carpeta.exists()


# ---------------------------------------------------------------------- fuente
def test_elegir_fuente_por_forma():
    from PIL import ImageFont
    disponibles = conversion.fuentes_disponibles()
    sans = next((f for f in disponibles if f.familia in ("Arial", "Liberation Sans")), None)
    serif = next((f for f in disponibles if f.familia in ("Times New Roman", "Liberation Serif")), None)
    if sans is None or serif is None:
        pytest.skip("Faltan fuentes para comparar")
    for elegida in (sans, serif):
        f = ImageFont.truetype(elegida.regular, 44, index=elegida.indice_regular)
        muestras = []
        for palabra in ("procedimiento", "laboratorio", "limpieza", "equipos", "analista", "Registrar"):
            render = conversion._render(f, palabra)
            muestras.append(conversion.Muestra(palabra, render, 0.1, 3.0, 4.0))
        mejor, alternativas = conversion.elegir_fuente(muestras, disponibles)
        assert mejor.clase == elegida.clase
        assert all(a.clase == elegida.clase for a in alternativas)


# ---------------------------------------------------------------------- de punta a punta (Tesseract)
def _similitud(a: str, b: str) -> float:
    return difflib.SequenceMatcher(None, a.lower(), b.lower()).ratio()


def _recuperadas(esperadas: list[str], encontradas: list[str]) -> float:
    falta = Counter(esperadas) - Counter(encontradas)
    return 1 - sum(falta.values()) / max(1, len(esperadas))


@pytest.fixture(scope="module")
def sop_convertido(tmp_path_factory):
    if not _hay_tesseract():
        pytest.skip("Tesseract con español no está instalado")
    carpeta = tmp_path_factory.mktemp("sop")
    original, escaneado = carpeta / "sop.pdf", carpeta / "sop_escaneado.pdf"
    esperado = sop.crear_sop(original)
    sop.escanear(original, escaneado)
    huella = hash_archivo(escaneado)
    exe, tessdata = conversion.herramientas_ocr()
    avance = []
    rec = conversion.convertir(conversion.paginas_de_pdf(escaneado, None, None), exe, tessdata, carpeta / "trabajo",
                               2, lambda h, t: avance.append((h, t)))
    assert hash_archivo(escaneado) == huella                         # el escaneo no se toca
    return rec, esperado, carpeta, avance


@requiere_tesseract
def test_sop_estructura(sop_convertido):
    rec, esperado, _, avance = sop_convertido
    assert avance == [(1, 2), (2, 2)]
    assert len(rec.paginas) == 2 and rec.fuente is not None and rec.fuente.clase == "sans"
    assert abs(rec.tamano_cuerpo - sop.TAMANO_CUERPO) <= 0.5
    cuerpo = [b for p in rec.paginas for b in p.bloques]
    titulos = [texto_de(b) for b in cuerpo if b.tipo == Tipo.TITULO]
    for t in esperado.titulos:
        assert max(_similitud(t, x) for x in titulos) >= 0.9, t
    nivel1 = [b for b in cuerpo if b.tipo == Tipo.TITULO and b.nivel == 1]
    assert len(nivel1) == 1 and nivel1[0].alineacion == Alineacion.CENTRO
    vinetas = [texto_de(b) for b in cuerpo if b.tipo == Tipo.ITEM and not b.numerado]
    numerados = [b for b in cuerpo if b.tipo == Tipo.ITEM and b.numerado]
    assert len(vinetas) == len(esperado.vinetas) and len(numerados) == len(esperado.numerados)
    assert all(_similitud(a, b) >= 0.9 for a, b in zip(vinetas, esperado.vinetas))
    assert [b.numero for b in numerados] == [1, 2, 3, 4]
    parrafos = [texto_de(b) for b in cuerpo if b.tipo == Tipo.PARRAFO]
    for p in esperado.parrafos:                                     # renglones unidos en un solo párrafo
        assert max(_similitud(p, x) for x in parrafos) >= 0.95, p


@requiere_tesseract
def test_sop_encabezado_pie_y_margenes(sop_convertido):
    rec, esperado, _, _ = sop_convertido
    for pag in rec.paginas:
        encabezado = " ".join(texto_de(b) for b in pag.encabezado)
        assert sop.CODIGO in encabezado and "03" in encabezado
        assert any(b.tipo == Tipo.TABLA for b in pag.encabezado)
        assert re.search(r"Página \d de 2", " ".join(texto_de(b) for b in pag.pie))
        assert all(sop.CODIGO not in texto_de(b) for b in pag.bloques)
    izq, arriba, der, abajo = rec.margenes
    assert abs(izq - sop.MARGEN) < 6 and abs(der - sop.MARGEN) < 6
    assert 88 < arriba < 130 and 30 < abajo < 80
    assert abs(rec.paginas[0].ancho - sop.ANCHO) < 1 and abs(rec.paginas[0].alto - sop.ALTO) < 1


@requiere_tesseract
def test_sop_tablas_firmas_y_sello(sop_convertido):
    rec, esperado, _, _ = sop_convertido
    tablas = [b for p in rec.paginas for b in p.bloques if b.tipo == Tipo.TABLA]
    assert len(tablas) == 2
    for tabla, filas in zip(tablas, esperado.tablas):
        assert len(tabla.filas) - 1 == len(filas) and len(tabla.columnas) - 1 == len(filas[0])
        celdas = {(c.fila, c.columna): " ".join(texto_de(x) for x in c.bloques if x.tipo != Tipo.IMAGEN)
                  for c in tabla.celdas}
        for i, fila in enumerate(filas):
            for j, texto in enumerate(fila):
                if texto:
                    assert _similitud(celdas[(i, j)], texto) >= 0.85, (texto, celdas[(i, j)])
    encabezados = next(c for c in tablas[0].celdas if (c.fila, c.columna) == (0, 0)).bloques[0]
    assert encabezados.negrita
    firmas = [x for c in tablas[1].celdas if c.fila == 1 for x in c.bloques if x.tipo == Tipo.IMAGEN]
    assert len(firmas) == esperado.firmas
    sueltas = [b for b in rec.paginas[1].bloques if b.tipo == Tipo.IMAGEN]
    assert len(sueltas) == esperado.sellos and sueltas[0].clase == "sello"
    assert all(Path(b.archivo).is_file() for b in firmas + sueltas)


@requiere_tesseract
def test_sop_recupera_el_texto(sop_convertido):
    rec, esperado, _, _ = sop_convertido
    encontradas = [w for p in rec.paginas for b in p.bloques for w in texto_de(b).split()]
    assert _recuperadas(esperado.palabras(), encontradas) >= 0.95


@requiere_tesseract
def test_sop_a_docx(sop_convertido):
    from docx import Document
    rec, esperado, carpeta, _ = sop_convertido
    destino = carpeta / "sop.docx"
    assert exportar_docx(rec, destino) == []
    doc = Document(str(destino))
    assert len(doc.tables) == 2
    assert [len(t.rows) for t in doc.tables] == [len(t) for t in esperado.tablas]
    estilos = {p.style.name for p in doc.paragraphs}
    assert {"Heading 1", "Heading 2", "List Paragraph"} <= estilos
    texto = texto_docx(destino)
    assert _recuperadas(esperado.palabras(), texto.split()) >= 0.95
    assert sop.CODIGO in " ".join(p.text for t in doc.sections[0].header.tables for r in t.rows
                                  for p in r.cells[2].paragraphs)
    assert doc.sections[0].header.tables[0].cell(0, 0)._tc.xpath(".//w:drawing")    # el logo
    assert doc.inline_shapes and len(doc.inline_shapes) >= esperado.firmas + esperado.sellos


@requiere_tesseract
def test_sop_a_pdf_con_texto(sop_convertido):
    rec, esperado, carpeta, _ = sop_convertido
    destino = carpeta / "sop_texto.pdf"
    assert exportar_pdf(rec, destino) == []
    with pymupdf.open(destino) as doc:
        palabras = [w[4] for p in doc for w in p.get_text("words")]
    assert _recuperadas(esperado.palabras(), palabras) >= 0.95


@requiere_tesseract
def test_sop_con_serifa_elige_fuente_con_serifa(tmp_path):
    original, escaneado = tmp_path / "sop.pdf", tmp_path / "esc.pdf"
    sop.crear_sop(original, serif=True)
    sop.escanear(original, escaneado, angulo=-0.8, semilla=7)
    exe, tessdata = conversion.herramientas_ocr()
    rec = conversion.convertir(conversion.paginas_de_pdf(escaneado, None, [0]), exe, tessdata, tmp_path / "t")
    assert rec.fuente is not None and rec.fuente.clase == "serif"
    assert any(b.tipo == Tipo.ITEM for b in rec.paginas[0].bloques)


@requiere_tesseract
def test_sop_fotografiado(tmp_path):
    """Una foto de la primera página, inclinada y en perspectiva, pasa por la importación de fotos."""
    from pdf_studio.core.escaneo import AjusteFoto, OpcionesEscaneo
    from tests import fotos_sinteticas as fs
    original = tmp_path / "sop.pdf"
    esperado = sop.crear_sop(original)
    with pymupdf.open(original) as doc:
        pix = doc[0].get_pixmap(dpi=230)
    papel = np.frombuffer(pix.samples, np.uint8).reshape(pix.height, pix.width, pix.n)[:, :, :3][:, :, ::-1].copy()
    img, _ = fs.foto(angulo=5, perspectiva=0.04, lienzo=(2300, 2900), papel=papel, sombra=True)
    foto = tmp_path / "foto.jpg"
    fs.guardar(img, foto)
    exe, tessdata = conversion.herramientas_ocr()
    rec = conversion.convertir(conversion.paginas_de_fotos([AjusteFoto(str(foto))], OpcionesEscaneo()), exe,
                               tessdata, tmp_path / "t")
    encontradas = [w for p in rec.paginas for b in p.todos() for w in texto_de(b).split()]
    primera = [w for t in esperado.titulos[:5] + esperado.parrafos[:2] + esperado.vinetas for w in t.split()]
    assert _recuperadas(primera, encontradas) >= 0.85
    assert abs(rec.paginas[0].alto / rec.paginas[0].ancho - sop.ALTO / sop.ANCHO) < 0.03


# ---------------------------------------------------------------------- interfaz
def _rec_con_dudosas(tmp_path) -> Reconstruccion:
    from PIL import Image
    h = Hoja()
    h.renglon("El equipo se limpia todos los días", 72, 100)
    h.renglon("Firmó el supervisor del turno", 72, 113)
    for po in h.a.palabras:
        if po.palabra.texto in ("limpia", "supervisor"):
            po.palabra.conf = 42.0
    imagen = tmp_path / "pagina.jpg"
    Image.new("RGB", (1240, 1754), "white").save(imagen)
    h.a.imagen = str(imagen)
    return maquetar([h.a])


def test_revision_navega_dudosas_con_enter(qtbot, tmp_path):
    from PySide6.QtCore import Qt

    from pdf_studio.ui.conversion import DialogoRevision
    rec = _rec_con_dudosas(tmp_path)
    dialogo = DialogoRevision(rec)
    qtbot.addWidget(dialogo)
    dialogo.show()
    assert [rec.buscar(i).texto for i in dialogo.dudosas] == ["limpia", "supervisor"]
    assert dialogo.campo.text() == "limpia"                       # arranca en la primera dudosa
    assert "limpia" in dialogo.resultado.toHtml()
    qtbot.keyClicks(dialogo.campo, "lava")                         # el texto estaba seleccionado: se reemplaza
    qtbot.keyClick(dialogo.campo, Qt.Key.Key_Return)
    palabra = next(p for p in rec.palabras() if p.texto == "lava")
    assert palabra.revisada and not palabra.dudosa
    assert dialogo.campo.text() == "supervisor"                    # pasó a la siguiente
    qtbot.keyClick(dialogo.campo, Qt.Key.Key_Return)               # se confirma tal cual
    assert rec.dudosas() == []
    assert "quedan 0" in dialogo.etiqueta_dudosas.text()


def test_revision_clic_en_cualquier_palabra(qtbot, tmp_path):
    from PySide6.QtCore import QUrl

    from pdf_studio.ui.conversion import DialogoRevision
    rec = _rec_con_dudosas(tmp_path)
    dialogo = DialogoRevision(rec)
    qtbot.addWidget(dialogo)
    otra = next(p for p in rec.palabras() if p.texto == "equipo")
    dialogo._enlace(QUrl(f"w:{otra.id}"))
    assert dialogo.campo.text() == "equipo" and dialogo.actual == otra.id
    dialogo.campo.setText("aparato")
    dialogo.confirmar()
    assert otra.texto == "aparato"


def test_dialogo_origen(qtbot):
    from pdf_studio.ui.conversion import DialogoOrigenConversion, Origen
    d = DialogoOrigenConversion("sop.pdf", 3)
    qtbot.addWidget(d)
    assert d.origen == Origen.ABIERTO
    d.grupo.button(2).setChecked(True)
    assert d.origen == Origen.FOTOS
    sin = DialogoOrigenConversion(None)
    qtbot.addWidget(sin)
    assert sin.origen == Origen.PDF


def test_accion_en_herramientas_sin_documento(ventana):
    assert ventana.acc_convertir_editable.isEnabled()
    assert "editable" in ventana.acc_convertir_editable.text()


def test_ventana_sin_tesseract_avisa(ventana, monkeypatch):
    from pdf_studio.dependencias import chequeo
    monkeypatch.setattr(chequeo, "buscar_tesseract", lambda: None)
    monkeypatch.setattr(ventana, "chequear_dependencias", lambda mostrar=True: None)
    assert not ventana.convertir_editable()
    titulo, texto, error = ventana.mensajes[-1]
    assert "Tesseract" in texto and not error                       # limitación: ícono de información


@requiere_tesseract
def test_ventana_convierte_el_pdf_abierto(ventana, tmp_path, monkeypatch, qtbot):
    original, escaneado = tmp_path / "sop.pdf", tmp_path / "escaneado.pdf"
    sop.crear_sop(original)
    sop.escanear(original, escaneado)
    huella = hash_archivo(escaneado)
    assert ventana.abrir(escaneado)
    destino, destino_pdf = tmp_path / "salida.docx", tmp_path / "salida.pdf"
    revisadas = []
    monkeypatch.setattr(ventana, "_dialogo_origen_conversion", lambda: "abierto")
    monkeypatch.setattr(ventana, "_dialogo_revision", lambda rec: revisadas.append(rec) or (rec.fuente, True))
    monkeypatch.setattr(ventana, "_pedir_ruta_docx", lambda sugerida: str(destino))
    monkeypatch.setattr(ventana, "_pedir_ruta_guardar", lambda sugerida: str(destino_pdf))
    abiertos = []
    monkeypatch.setattr(ventana, "_preguntar_abrir_documento", lambda texto, avisos: abiertos.append(texto))
    assert ventana.convertir_editable()
    qtbot.waitUntil(lambda: bool(abiertos), timeout=300_000)
    assert destino.is_file() and destino_pdf.is_file()
    assert "Se creó «salida.docx» y «salida.pdf»." in abiertos[0] and "Atención" not in abiertos[0]
    assert sop.CODIGO in texto_docx(destino)
    assert hash_archivo(escaneado) == huella and not ventana.documento.modificado


@requiere_tesseract
def test_ventana_no_pisa_el_original(ventana, tmp_path, monkeypatch, qtbot):
    original, escaneado = tmp_path / "sop.pdf", tmp_path / "escaneado.pdf"
    sop.crear_sop(original)
    sop.escanear(original, escaneado)
    huella = hash_archivo(escaneado)
    monkeypatch.setattr(ventana, "_dialogo_origen_conversion", lambda: "pdf")
    monkeypatch.setattr(ventana, "_pedir_ruta_comparar", lambda titulo="": str(escaneado))
    monkeypatch.setattr(ventana, "_dialogo_revision", lambda rec: (rec.fuente, True))
    monkeypatch.setattr(ventana, "_pedir_ruta_docx", lambda sugerida: str(tmp_path / "x.docx"))
    monkeypatch.setattr(ventana, "_pedir_ruta_guardar", lambda sugerida: str(escaneado))
    assert ventana.convertir_editable()
    qtbot.waitUntil(lambda: any("reemplazar al original" in m[1] for m in ventana.mensajes), timeout=300_000)
    assert hash_archivo(escaneado) == huella and not (tmp_path / "x.docx").exists()
