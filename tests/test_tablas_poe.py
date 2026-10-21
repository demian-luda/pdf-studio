"""Tablas de SOPs / POEs viejos al convertir un escaneado: líneas finas, inclinadas, curvas o
cortadas; celdas combinadas; etiqueta y valor en la misma celda; tablas cortadas por el borde
de la hoja; sin texto duplicado debajo de imágenes; negritas relativas al tamaño; corrector
ortográfico. Al final, un POE sintético fotografiado (perspectiva, curvatura y reflejo) de
punta a punta con Tesseract."""
from __future__ import annotations

import difflib
from collections import Counter

import numpy as np
import pytest

pytest.importorskip("cv2")
pytest.importorskip("docx")

import cv2  # noqa: E402

from pdf_studio.core import conversion, maquetado, ortografia  # noqa: E402
from pdf_studio.core.exportar_editable import exportar_docx, texto_docx  # noqa: E402
from pdf_studio.core.maquetado import AnalisisPagina, Grilla, PalabraOCR, maquetar, texto_de  # noqa: E402
from pdf_studio.core.reconstruccion import Bloque, Linea, Palabra, Reconstruccion, Tipo  # noqa: E402
from tests import sops_sinteticos as sop  # noqa: E402
from tests.conftest import hash_archivo  # noqa: E402
from tests.test_conversion import Hoja, requiere_tesseract  # noqa: E402


def _lienzo(w: int = 1600, h: int = 1200) -> np.ndarray:
    return np.zeros((h, w), np.uint8)


def _tabla(tinta: np.ndarray, ys: list[int], xs: list[int], grosor: int = 3, inclinacion: float = 0.0) -> None:
    """Dibuja la grilla; con `inclinacion` (px por px) las líneas quedan torcidas como en una foto."""
    for y in ys:
        cv2.line(tinta, (xs[0], y), (xs[-1], int(y + inclinacion * (xs[-1] - xs[0]))), 255, grosor)
    for x in xs:
        cv2.line(tinta, (x, ys[0]), (int(x - inclinacion * (ys[-1] - ys[0])), ys[-1]), 255, grosor)


# ---------------------------------------------------------------------- detección de la grilla
def test_lineas_largas_inclinadas_siguen_siendo_finas():
    """Una línea larga apenas inclinada tiene una caja alta, pero es fina: no se descarta."""
    tinta = _lienzo(3000, 2400)
    ys = [300 + 120 * i for i in range(9)]
    _tabla(tinta, ys, [200, 1000, 1600, 2200, 2800], inclinacion=0.012)
    grillas, _ = conversion.detectar_grillas(tinta)
    assert len(grillas) == 1
    g = grillas[0]
    assert (len(g.filas) - 1, len(g.columnas) - 1) == (8, 4) and len(g.celdas) == 32


def test_lineas_cortadas_se_completan_con_las_filas_repetidas():
    tinta = _lienzo()
    ys, xs = [100 + 80 * i for i in range(7)], [100, 500, 900, 1300]
    _tabla(tinta, ys, xs)
    cv2.line(tinta, (500, 355), (500, 405), 0, 9)              # vertical cortada en una fila (queda un rastro)
    cv2.line(tinta, (140, 420), (460, 420), 0, 9)              # horizontal casi borrada en una columna
    grillas, _ = conversion.detectar_grillas(tinta)
    assert len(grillas) == 1
    assert len(grillas[0].celdas) == 18                         # 6 × 3, ninguna combinada


def test_celda_combinada_de_verdad_se_respeta():
    """En una tabla regular, una celda combinada (sin ningún rastro de línea y con las otras
    líneas de esa columna también faltando) no se parte."""
    tinta = _lienzo()
    ys, xs = [100 + 80 * i for i in range(6)], [100, 500, 900, 1300]
    _tabla(tinta, ys, xs)
    cv2.line(tinta, (104, 260), (496, 260), 0, 9)              # la primera columna combina dos filas
    grillas, _ = conversion.detectar_grillas(tinta)
    assert (1, 0, 2, 1) in grillas[0].celdas and len(grillas[0].celdas) == 14


def test_tabla_de_un_renglon_cortada_por_el_borde_de_la_hoja():
    """Verticales cortas (una sola fila) y el borde de abajo fuera de la hoja."""
    tinta = _lienzo(2400, 1000)
    cv2.line(tinta, (200, 930), (2200, 930), 255, 4)
    for x in (200, 1400, 2200):
        cv2.line(tinta, (x, 930), (x, 999), 255, 4)
    grillas, _ = conversion.detectar_grillas(tinta)
    assert len(grillas) == 1
    g = grillas[0]
    assert (len(g.filas) - 1, len(g.columnas) - 1) == (1, 2)


def test_trazos_de_letras_grandes_no_son_lineas():
    tinta = _lienzo(2400, 1800)
    _tabla(tinta, [200, 500], [200, 1200, 2200])
    for x in range(300, 700, 70):                               # «letras» de 60 px dentro de la celda
        cv2.line(tinta, (x, 300), (x, 360), 255, 9)
    grillas, mascara = conversion.detectar_grillas(tinta)
    assert len(grillas[0].columnas) == 3
    assert mascara[330, 300] == 0                               # no se borran antes del OCR


# ---------------------------------------------------------------------- etiqueta / valor
def _celdas_de(bloque: Bloque) -> dict[tuple[int, int], str]:
    return {(c.fila, c.columna): " ".join(texto_de(b) for b in c.bloques if b.tipo != Tipo.IMAGEN).strip()
            for c in bloque.celdas}


def test_etiqueta_y_valor_van_en_celdas_propias():
    h = Hoja()
    filas = [100.0, 120.0, 140.0, 160.0]
    h.a.grillas.append(Grilla(filas, [60.0, 200.0, 540.0], [(0, 0, 3, 1)] + [(f, 1, 1, 1) for f in range(3)]))
    h.renglon("CONTROL", 90, 114, negrita=True)
    for f, (etiqueta, valor) in enumerate([("POE N°:", "POE-CC-99-01"), ("Vigente desde:", "Mayo 2024"),
                                           ("Fecha: 14/05/2024", None)]):
        h.renglon(etiqueta, 205, filas[f] + 14)
        if valor:
            h.renglon(valor, 340, filas[f] + 14, negrita=True)
    tabla = next(b for b in maquetar([h.a]).paginas[0].todos() if b.tipo == Tipo.TABLA)
    assert len(tabla.columnas) == 4 and tabla.columnas_sin_borde == [2]
    celdas = _celdas_de(tabla)
    assert celdas[(0, 1)] == "POE N°:" and celdas[(0, 2)] == "POE-CC-99-01"
    assert celdas[(1, 1)] == "Vigente desde:" and celdas[(1, 2)] == "Mayo 2024"
    assert celdas[(2, 1)] == "Fecha: 14/05/2024" and celdas[(2, 2)] == ""
    assert celdas[(0, 0)] == "CONTROL" and any(c.filas == 3 for c in tabla.celdas)


def test_docx_no_dibuja_el_borde_agregado(tmp_path):
    from docx import Document
    from docx.oxml.ns import qn
    h = Hoja()
    h.a.grillas.append(Grilla([100.0, 120.0], [60.0, 540.0], [(0, 0, 1, 1)]))
    h.renglon("Preparado por:", 66, 114)
    h.renglon("ANA GÓMEZ", 300, 114)
    rec = maquetar([h.a])
    assert exportar_docx(rec, tmp_path / "t.docx") == []
    tabla = Document(tmp_path / "t.docx").tables[0]
    assert len(tabla.columns) == 2
    assert [c.text for c in tabla.rows[0].cells] == ["Preparado por:", "ANA GÓMEZ"]
    derecha = tabla.cell(0, 0)._tc.tcPr.find(qn("w:tcBorders")).find(qn("w:right"))
    assert derecha.get(qn("w:val")) == "nil"


# ---------------------------------------------------------------------- imágenes y texto
def test_region_que_es_texto_legible_no_va_como_imagen():
    tinta = _lienzo(1200, 400)
    cv2.putText(tinta, "Firma", (100, 200), cv2.FONT_HERSHEY_SIMPLEX, 2.0, 255, 4)
    region = conversion._Region((90, 140, 320, 215), "imagen", False)
    palabras = [conversion.PalabraTSV("Firma", (95, 145, 315, 212), 58, (1, 1, 1))]
    assert conversion._es_texto(region, tinta, palabras, 60.0)
    # Un garabato al lado del texto sí es imagen
    xs = np.linspace(330, 700, 60)
    cv2.polylines(tinta, [np.stack([xs, 180 + 30 * np.sin(xs / 15)], 1).astype(np.int32)], False, 255, 3)
    region = conversion._Region((90, 140, 710, 215), "imagen", False)
    assert not conversion._es_texto(region, tinta, palabras, 60.0)


def test_sombra_del_borde_de_la_hoja():
    tinta = _lienzo(800, 600)
    tinta[580:, :] = 255                                         # sombra abajo
    cv2.putText(tinta, "N/A", (300, 560), cv2.FONT_HERSHEY_SIMPLEX, 1.5, 255, 3)
    bordes = conversion.tinta_en_bordes(tinta)
    assert bordes[590, 400] and not bordes[545, 330]


# ---------------------------------------------------------------------- negritas
def test_negrita_relativa_al_tamano_de_letra():
    """Un título grande en regular tiene el trazo más grueso que el cuerpo, pero no es negrita;
    un rótulo chico en negrita tiene el trazo más fino que el título, pero sí lo es."""
    regular, negrita, exceso = 0.075, 0.115, 0.35

    def linea(texto: str, tamano: float, en_negrita: bool) -> Linea:
        grosor = (negrita if en_negrita else regular) * tamano + exceso
        return Linea([Palabra(f"x{i}", t, (0, 0, 1, 1), 95, trazo=grosor) for i, t in enumerate(texto.split())],
                     tamano)

    cuerpo = [linea("Revisado por el supervisor del turno", 10, False) for _ in range(12)]
    titulo = linea("Uso de las estufas", 17, False)
    rotulo = linea("Actualizaciones vigentes", 9, True)
    valor = linea("Mayo 2024", 10, True)
    maquetado.marcar_negritas(cuerpo + [titulo, rotulo, valor], (regular, negrita))
    assert not titulo.negrita and rotulo.negrita and valor.negrita
    assert not any(l.negrita for l in cuerpo)


def test_mayuscula_confundida_por_el_ocr():
    """«vence:» leído en minúscula da un tamaño enorme: es «Vence:» y queda para revisar."""
    h = Hoja()
    for i in range(6):
        h.renglon("Revisado por el supervisor", 72, 100 + 20 * i)
    palabra = Palabra("v", "vence:", (300, 92.8, 330, 100), 95)          # alto de una mayúscula de 10 pt
    h.a.palabras.append(PalabraOCR(palabra, (99, 0, 0)))
    rec = maquetar([h.a])
    assert palabra.texto == "Vence:" and palabra.dudosa


# ---------------------------------------------------------------------- ortografía
def test_corrector_marca_lo_que_no_existe():
    s = ortografia.sospechosa
    assert s("Fima:") and s("por-") and s("Vicro»io") and s("Arobado")
    assert not any(s(p) for p in ("Firma:", "por:", "Microbiología", "estufas", "Vence:", "reemplazado:"))
    # Números, códigos, siglas y nombres propios no se marcan
    assert not any(s(p) for p in ("14/05/2024", "POE-XX-00-01", "N°:", "N/A", "POE", "SOP", "RIVAROLA",
                                  "Gutierrez", "Lucía"))
    # Una palabra partida al final del renglón no es un error
    assert not s("micro-", "biología") and s("micro-")


def test_sospechosas_se_marcan_como_dudosas_aunque_el_ocr_confie():
    h = Hoja()
    h.renglon("Fima del responsable", 72, 100, conf=96)
    rec = maquetar([h.a])
    assert ortografia.marcar_sospechosas(rec) == 1
    assert [p.texto for p in rec.dudosas()] == ["Fima"]
    p = rec.dudosas()[0]
    p.texto, p.revisada = "Firma", True
    assert rec.dudosas() == []
    assert Reconstruccion.desde_dict(rec.a_dict()).paginas[0].bloques[0].lineas[0].palabras[0].sospechosa


# ---------------------------------------------------------------------- POE fotografiado
def _parecido(a: str, b: str) -> float:
    return difflib.SequenceMatcher(None, a.lower(), b.lower()).ratio()


@pytest.fixture(scope="module")
def poe_convertido(tmp_path_factory):
    if not conversion_disponible():
        pytest.skip("Tesseract con español no está instalado")
    carpeta = tmp_path_factory.mktemp("poe")
    original, foto = carpeta / "poe.pdf", carpeta / "poe_foto.pdf"
    esperado = sop.crear_poe(original)
    sop.fotografiar(original, foto)
    huella = hash_archivo(foto)
    exe, tessdata = conversion.herramientas_ocr()
    rec = conversion.convertir(conversion.paginas_de_pdf(foto, None, None), exe, tessdata, carpeta / "trabajo", 1)
    assert hash_archivo(foto) == huella
    return rec, esperado, carpeta


def conversion_disponible() -> bool:
    try:
        conversion.herramientas_ocr()
        return True
    except Exception:
        return False


def _tablas(rec: Reconstruccion) -> list[Bloque]:
    return [b for pag in rec.paginas for b in pag.todos() if b.tipo == Tipo.TABLA]


def _forma(b: Bloque) -> tuple[int, int]:
    return len(b.filas) - 1, len(b.columnas) - 1


@requiere_tesseract
def test_poe_cantidad_y_forma_de_las_tablas(poe_convertido):
    rec, _, _ = poe_convertido
    tablas = _tablas(rec)
    assert [_forma(t) for t in tablas] == [(3, 3), (1, 1), (3, 4), (8, 4), (1, 4)]
    encabezado, _, aprobacion, actualizaciones, reemplaza = tablas
    # La celda «CONTROL DE CALIDAD» combina las tres filas; el resto no combina nada
    assert [(c.fila, c.columna, c.filas, c.columnas) for c in encabezado.celdas if c.filas > 1] == [(0, 0, 3, 1)]
    for t in (aprobacion, actualizaciones, reemplaza):
        assert len(t.celdas) == _forma(t)[0] * _forma(t)[1]
    # Las columnas de los valores no tenían línea: se agregan sin borde
    assert encabezado.columnas_sin_borde == [2]
    assert aprobacion.columnas_sin_borde == [1] and reemplaza.columnas_sin_borde == [1, 3]
    assert actualizaciones.columnas_sin_borde == []


@requiere_tesseract
def test_poe_cada_texto_en_su_celda(poe_convertido):
    rec, esperado, _ = poe_convertido
    encabezado, titulo, aprobacion, actualizaciones, reemplaza = _tablas(rec)
    celdas = _celdas_de(encabezado)
    assert _parecido(celdas[(0, 0)], esperado.combinada) >= 0.9
    for f, (etiqueta, valor) in enumerate(esperado.encabezado):
        assert _parecido(celdas[(f, 1)], etiqueta) >= 0.8, (f, celdas[(f, 1)])
        assert _parecido(celdas[(f, 2)], valor) >= 0.8, (f, celdas[(f, 2)])
    assert _parecido(_celdas_de(titulo)[(0, 0)], esperado.titulo) >= 0.9
    for tabla, filas in ((aprobacion, esperado.aprobacion), (actualizaciones, esperado.actualizaciones),
                         (reemplaza, esperado.reemplaza)):
        celdas = _celdas_de(tabla)
        for f, fila in enumerate(filas):
            for c, texto in enumerate(fila):
                assert _parecido(celdas[(f, c)], texto) >= 0.75, (f, c, celdas[(f, c)], texto)


@requiere_tesseract
def test_poe_negritas_y_titulos(poe_convertido):
    rec, esperado, _ = poe_convertido
    encabezado = _tablas(rec)[0]
    valores = [b for c in encabezado.celdas if c.columna == 2 for b in c.bloques if b.lineas]
    assert valores and all(b.negrita for b in valores)
    etiquetas = [b for c in encabezado.celdas if c.columna == 1 for b in c.bloques if b.lineas]
    assert etiquetas and not any(b.negrita for b in etiquetas)
    secciones = [texto_de(b) for pag in rec.paginas for b in pag.bloques if b.tipo == Tipo.TITULO]
    for s in esperado.secciones:
        assert any(_parecido(s, t) >= 0.85 for t in secciones), s


@requiere_tesseract
def test_poe_sin_duplicados_y_firma_como_imagen(poe_convertido):
    rec, esperado, carpeta = poe_convertido
    aprobacion = _tablas(rec)[2]
    firma = next(c for c in aprobacion.celdas if (c.fila, c.columna) == (0, 3))
    assert [b.clase for b in firma.bloques if b.tipo == Tipo.IMAGEN] == ["firma"]
    imagenes = [b for pag in rec.paginas for b in pag.todos() if b.tipo == Tipo.IMAGEN]
    imagenes += [x for t in _tablas(rec) for c in t.celdas for x in c.bloques if x.tipo == Tipo.IMAGEN]
    assert sorted(b.clase for b in imagenes) == ["firma", "logo"]
    from tests.robustez_pdfs_reales import estructura_de_tablas
    assert estructura_de_tablas(rec).endswith("palabras bajo imágenes 0")
    # En el .docx no sobra ninguna palabra (cada una aparece a lo sumo tantas veces como en el POE)
    assert exportar_docx(rec, carpeta / "poe.docx") == []
    encontradas = Counter(texto_docx(carpeta / "poe.docx").split())
    esperadas = Counter(esperado.palabras())
    for palabra in ("Firma:", "Fecha:", "Vence:", "Revisado", "N/A", "Reemplaza", "CONTROL"):
        assert encontradas[palabra] <= esperadas[palabra], palabra
    total = sum(encontradas.values())
    assert total <= sum(esperadas.values()) + 3          # el logo puede sumar alguna palabra suelta


@requiere_tesseract
def test_poe_docx_con_las_tablas_de_word(poe_convertido):
    from docx import Document
    rec, _, carpeta = poe_convertido
    destino = carpeta / "poe_tablas.docx"
    assert exportar_docx(rec, destino) == []
    doc = Document(destino)
    tablas = list(doc.tables) + [t for s in doc.sections for t in s.header.tables]
    formas = sorted((len(t.rows), len(t.columns)) for t in tablas)
    assert formas == sorted([(3, 3), (1, 1), (3, 4), (8, 4), (1, 4)])
